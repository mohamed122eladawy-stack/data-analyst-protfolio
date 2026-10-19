"""Builds an interactive Excel sales dashboard (all numbers are live formulas)."""
import random
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

random.seed(21)
BRANCHES = ["Cairo", "Alexandria", "Giza", "Mansoura"]
PRODUCTS = {"Laptop": ("Electronics", 18500, 15000), "Monitor": ("Electronics", 4200, 3300),
            "Printer": ("Electronics", 3800, 3000), "Keyboard": ("Accessories", 450, 280),
            "Mouse": ("Accessories", 250, 140), "Router": ("Networking", 900, 650),
            "Webcam": ("Accessories", 700, 480)}
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

F = "Arial"
NAVY = PatternFill("solid", fgColor="1F3864")
LIGHT = PatternFill("solid", fgColor="DDEBF7")
YELLOW = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ds = wb.active
ds.title = "Data"
dash = wb.create_sheet("Dashboard", 0)

# ---------- Data sheet ----------
heads = ["Date", "Month", "Branch", "Product", "Category", "Qty", "Unit Price", "Unit Cost", "Revenue", "Profit"]
ds.append(heads)
N = 600
for r in range(2, N + 2):
    p = random.choice(list(PRODUCTS))
    cat, price, cost = PRODUCTS[p]
    d = date(2026, 1, 1) + timedelta(days=random.randint(0, 364))
    ds.append([d, f"=MONTH(A{r})", random.choice(BRANCHES), p, cat, random.randint(1, 12), price, cost,
               f"=F{r}*G{r}", f"=F{r}*(G{r}-H{r})"])
    ds[f"A{r}"].number_format = "yyyy-mm-dd"
    for c in "GHIJ":
        ds[f"{c}{r}"].number_format = "#,##0"
for c in ds[1]:
    c.font = Font(name=F, bold=True, color="FFFFFF")
    c.fill = NAVY
    c.alignment = Alignment(horizontal="center")
for row in ds.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, size=10)
for col, w in zip("ABCDEFGHIJ", [12, 8, 13, 13, 14, 7, 12, 12, 13, 13]):
    ds.column_dimensions[col].width = w
ds.freeze_panes = "A2"
ds.auto_filter.ref = f"A1:J{N+1}"
last = N + 1
rng = lambda col: f"Data!${col}$2:${col}${last}"

# ---------- Dashboard ----------
dash.sheet_view.showGridLines = False
dash["A1"] = "Sales Performance Dashboard - 2026"
dash["A1"].font = Font(name=F, size=18, bold=True, color="1F3864")
dash["A2"] = "Select a branch in B4. All figures and charts update automatically. (Sample data - fictional)"
dash["A2"].font = Font(name=F, size=9, italic=True, color="7F7F7F")

dash["A4"], dash["B4"] = "Branch filter:", "All"
dash["A4"].font = Font(name=F, bold=True)
dash["B4"].fill = YELLOW
dash["B4"].font = Font(name=F, bold=True, color="0000FF")
dash["B4"].border = BOX
dv = DataValidation(type="list", formula1='"All,Cairo,Alexandria,Giza,Mansoura"', allow_blank=False)
dash.add_data_validation(dv)
dv.add("B4")
dash["C4"] = '=IF($B$4="All","*",$B$4)'     # helper criteria (wildcard when All)
dash["C4"].font = Font(name=F, size=8, color="A6A6A6")
dash["D4"] = "<- helper cell (do not edit)"
dash["D4"].font = Font(name=F, size=8, color="A6A6A6")

# KPI cards
kpis = [("Total Revenue", f'=SUMIFS({rng("I")},{rng("C")},$C$4)', "#,##0"),
        ("Total Profit", f'=SUMIFS({rng("J")},{rng("C")},$C$4)', "#,##0"),
        ("Profit Margin", "=IF(A7=0,0,B7/A7)", "0.0%"),
        ("Units Sold", f'=SUMIFS({rng("F")},{rng("C")},$C$4)', "#,##0"),
        ("Transactions", f'=COUNTIFS({rng("C")},$C$4)', "#,##0")]
kpi_cols = ["A", "B", "C", "D", "E"]
for col, (label, formula, fmt) in zip(kpi_cols, kpis):
    dash[f"{col}6"] = label
    dash[f"{col}6"].font = Font(name=F, size=9, bold=True, color="FFFFFF")
    dash[f"{col}6"].fill = NAVY
    dash[f"{col}6"].alignment = Alignment(horizontal="center")
    dash[f"{col}7"] = formula
    dash[f"{col}7"].number_format = fmt
    dash[f"{col}7"].font = Font(name=F, size=14, bold=True, color="1F3864")
    dash[f"{col}7"].fill = LIGHT
    dash[f"{col}7"].alignment = Alignment(horizontal="center")
dash.row_dimensions[7].height = 28

def table(top, headers, labels, formulas, fmts):
    for j, h in enumerate(headers):
        c = dash.cell(row=top, column=1 + j, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = NAVY
        c.alignment = Alignment(horizontal="center")
    for i, lab in enumerate(labels):
        r = top + 1 + i
        dash.cell(row=r, column=1, value=lab).font = Font(name=F)
        for j, (fx, fmt) in enumerate(zip(formulas, fmts)):
            c = dash.cell(row=r, column=2 + j, value=fx(r))
            c.number_format = fmt
            c.font = Font(name=F)
        for j in range(len(headers)):
            dash.cell(row=r, column=1 + j).border = BOX
    return top + 1, top + len(labels)

# Monthly table (filtered by branch)
m_top = 10
dash.cell(row=m_top - 1, column=1, value="Monthly Trend").font = Font(name=F, bold=True, size=12, color="1F3864")
m1, m2 = table(m_top, ["Month", "Revenue", "Profit"], MONTHS,
               [lambda r: f'=SUMIFS({rng("I")},{rng("B")},ROW()-{m_top},{rng("C")},$C$4)',
                lambda r: f'=SUMIFS({rng("J")},{rng("B")},ROW()-{m_top},{rng("C")},$C$4)'],
               ["#,##0", "#,##0"])

# Product table
p_top = m2 + 3
dash.cell(row=p_top - 1, column=1, value="Sales by Product").font = Font(name=F, bold=True, size=12, color="1F3864")
prods = list(PRODUCTS)
p1, p2 = table(p_top, ["Product", "Revenue", "Profit", "Units", "Share of Revenue"], prods,
               [lambda r: f'=SUMIFS({rng("I")},{rng("D")},$A{r},{rng("C")},$C$4)',
                lambda r: f'=SUMIFS({rng("J")},{rng("D")},$A{r},{rng("C")},$C$4)',
                lambda r: f'=SUMIFS({rng("F")},{rng("D")},$A{r},{rng("C")},$C$4)',
                lambda r: f"=IF($A$7=0,0,B{r}/$A$7)"],
               ["#,##0", "#,##0", "#,##0", "0.0%"])
dash.conditional_formatting.add(f"B{p1}:B{p2}", DataBarRule(start_type="min", end_type="max", color="5B9BD5"))

# Branch table (always shows all branches for comparison)
b_top = p2 + 3
dash.cell(row=b_top - 1, column=1, value="Branch Comparison (not affected by filter)").font = Font(name=F, bold=True, size=12, color="1F3864")
b1, b2 = table(b_top, ["Branch", "Revenue", "Profit", "Margin"], BRANCHES,
               [lambda r: f'=SUMIFS({rng("I")},{rng("C")},$A{r})',
                lambda r: f'=SUMIFS({rng("J")},{rng("C")},$A{r})',
                lambda r: f"=IF(B{r}=0,0,C{r}/B{r})"],
               ["#,##0", "#,##0", "0.0%"])
dash.cell(row=b2 + 1, column=1, value="Best branch:").font = Font(name=F, bold=True)
dash.cell(row=b2 + 1, column=2, value=f"=INDEX(A{b1}:A{b2},MATCH(MAX(B{b1}:B{b2}),B{b1}:B{b2},0))").font = Font(name=F, bold=True, color="C00000")

for col, w in zip("ABCDEFGHIJKLMN", [18, 15, 15, 14, 18, 3, 3, 3, 3, 3, 3, 3, 3, 3]):
    dash.column_dimensions[col].width = w

# Charts
line = LineChart()
line.title, line.height, line.width = "Monthly Revenue & Profit", 8, 16
line.add_data(Reference(dash, min_col=2, max_col=3, min_row=m_top, max_row=m2), titles_from_data=True)
line.set_categories(Reference(dash, min_col=1, min_row=m1, max_row=m2))
dash.add_chart(line, "G4")

bar = BarChart()
bar.type, bar.title, bar.height, bar.width = "bar", "Revenue by Product", 8, 16
bar.add_data(Reference(dash, min_col=2, min_row=p_top, max_row=p2), titles_from_data=True)
bar.set_categories(Reference(dash, min_col=1, min_row=p1, max_row=p2))
bar.legend = None
dash.add_chart(bar, "G21")

pie = PieChart()
pie.title, pie.height, pie.width = "Revenue Share by Branch", 8, 16
pie.add_data(Reference(dash, min_col=2, min_row=b_top, max_row=b2), titles_from_data=True)
pie.set_categories(Reference(dash, min_col=1, min_row=b1, max_row=b2))
dash.add_chart(pie, "G38")

wb.save("Sales_Dashboard.xlsx")
print("saved")
